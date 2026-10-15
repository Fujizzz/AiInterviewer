"""Regression tests for the round-3 data-quality fixes (residual markup, d2l/mkdocs/
docusaurus handling, junk-filter false positives, residue scanner, loaders)."""

import pytest

# Optional rag-extra dependencies: skip the whole module when they are not installed.
pytest.importorskip("openpyxl")
pytest.importorskip("requests")
import json

import openpyxl

from rag import cli, residue
from rag.clean import clean_markdown, count_fences, junk_category
from rag.htmlmd import convert_embedded_html, html_to_markdown
from rag.loaders.crawler_loader import Crawler
from rag.loaders.git_loader import substitute_snippet_includes_in_text
from rag.loaders.xlsx_loader import fetch as xlsx_fetch
from rag.loaders.xlsx_loader import parse_workbook
from rag.profiles import (
    profile_d2l,
    profile_markdown,
    profile_rst,
    resolve_tab_selected,
    rst_to_markdown,
)
from rag.registry import Source
from tests.rag.test_crawler_loader import PREFIX, FakeResp, Net, page

# ---------------------------------------------------------------------------
# d2l: tab.selected + dangling :numref: labels
# ---------------------------------------------------------------------------


def test_tab_selected_drops_other_frameworks_and_unwraps_pytorch():
    code = [
        "x = 1",
        "if tab.selected('mxnet'):",
        "    import mxnet as mx",
        "if tab.selected('tensorflow'):",
        "    import tensorflow as tf",
        "if tab.selected('jax'):",
        "    import jax",
        "if tab.selected('pytorch'):",
        "    import torch",
        "    for i in range(2):",
        "        print(i)",
        "y = 2",
    ]
    out = resolve_tab_selected(code)
    assert out == ["x = 1", "import torch", "for i in range(2):", "    print(i)", "y = 2"]


def test_tab_selected_or_combination_and_elif():
    out = resolve_tab_selected(
        [
            "if tab.selected('mxnet') or tab.selected('pytorch'):",
            "    a = 1",
            "elif tab.selected('tensorflow'):",
            "    a = 2",
        ]
    )
    assert out == ["a = 1"]


def test_d2l_cell_has_no_tab_selected_left():
    src = (
        "```{.python .input}\n#@tab all\nif tab.selected('mxnet'):\n    m = 1\nif "
        "tab.selected('pytorch'):\n    m = 2\n```\n"
    )
    out = profile_d2l(src)
    assert (
        "tab.selected" not in out
        and "m = 2" in out
        and "m = 1" not in out
        and "\n    m = 2" not in out
    )


def test_d2l_numref_becomes_readable_text_without_doubled_kind_word():
    out = profile_d2l(
        "As shown in :numref:`sec_resnet` and :numref:`fig_blocks_a`; see Section "
        ":numref:`sec_lstm`, "
        "Figure :numref:`fig_seq2seq`, Equation :eqref:`eq_x`.\n"
    )
    assert ":numref:" not in out and ":eqref:" not in out
    assert "the resnet section" in out and "the blocks a figure" in out
    assert "Section the lstm section" not in out and "Figure the seq2seq figure" not in out
    assert "sec_" not in out and "fig_" not in out and "Equation the equation" not in out


# ---------------------------------------------------------------------------
# mkdocs / pymdownx snippets
# ---------------------------------------------------------------------------


def test_snippet_marker_lines_are_stripped_everywhere():
    text = (
        "# T\n\n--8<-- [start:install]\nInstall it.\n--8<-- [end:install]\n\n    --8<-- [start:x]\n"
        "    more\n    --8<-- [end:x]\n"
    )
    out = profile_markdown(text)
    assert "--8<--" not in out and "Install it." in out and "more" in out


def test_marker_only_chunk_is_junk():
    assert junk_category("--8<-- [start:a]\n--8<-- [end:a]\n") == "doc-tool-directive-only"


def test_oversized_snippet_replaced_by_one_line_note_not_empty_fence():
    text = 'Intro.\n\n```python\n--8<-- "pkg/big.py"\n```\n\nAfter.'
    out, _, _, too_large = substitute_snippet_includes_in_text(
        text, {"pkg/big.py": "x = 1\n" * 5000}, max_tokens=100, permalink_base="https://h/blob/abc"
    )
    assert too_large == 1
    assert "(full content omitted for size; see upstream: https://h/blob/abc/pkg/big.py)" in out
    assert "```" not in out and "x = 1" not in out


def test_virtual_generated_include_is_not_given_a_bogus_permalink():
    out, _, not_found, _ = substitute_snippet_includes_in_text(
        '--8<-- "gen:engine-args"', {}, permalink_base="https://h/blob/abc"
    )
    assert not_found == 1 and "https://h" not in out and "build time" in out


def test_chunk_with_only_an_omission_note_is_junk():
    note = "(full content omitted for size; see upstream: https://h/blob/abc/pkg/big.py)"
    assert junk_category(note) == "doc-tool-directive-only"


def test_empty_fence_left_by_stripping_is_removed():
    out = profile_markdown("Text\n\n```python\n\n```\n\nMore")
    assert "```" not in out and "Text" in out and "More" in out


# ---------------------------------------------------------------------------
# leftover markup
# ---------------------------------------------------------------------------


def test_docusaurus_admonitions_with_titles_become_bold_label():
    out = profile_markdown(
        ":::tip Expert tip\n\nGet paid more.\n\n:::\n\n:::info\nplain\n:::\n\nafter\n"
    )
    assert ":::" not in out
    assert (
        "**Expert tip:**" in out
        and "Get paid more." in out
        and "**Info:**" in out
        and "after" in out
    )


def test_stray_triple_colon_closer_is_dropped():
    assert ":::" not in profile_markdown("text\n\n:::\n\nmore")


def test_mkdocstrings_directive_is_not_mistaken_for_admonition_and_is_stripped():
    out = profile_markdown(
        "Intro\n\n::: fastapi.FastAPI\n    options:\n      members:\n        - x\n\n## Next\n"
    )
    assert "fastapi.FastAPI" not in out and "options:" not in out and "## Next" in out


def test_heading_attribute_lists_are_stripped():
    out = profile_markdown(
        '## Title { #my-id }\n\n## T2 { data-toc-label="abc" }\n\n## T3 {: #kramdown-id }\n\n## '
        "`f{x}` code\n"
    )
    assert "## Title\n" in out and "## T2\n" in out and "## T3\n" in out
    assert "{ #" not in out and "data-toc" not in out and "{:" not in out
    assert "`f{x}`" in out  # braces that are not attribute lists survive


def test_image_attribute_list_is_stripped():
    out = profile_markdown(
        '![](./a.png){ align="center" alt="x" width="60%" } caption\n\nf(x){y}\n'
    )
    assert "{ align" not in out and "f(x){y}" in out


def test_language_switcher_line_is_removed():
    out = profile_markdown(
        "*[English](README.md) ∙ [日本語](README-ja.md) ∙ [简体中文](README-zh-Hans.md)*\n\n"
        "# Title\n"
    )
    assert "日本語" not in out and "# Title" in out


def test_html_img_keeps_meaningful_alt_only():
    out = convert_embedded_html(
        '<img src="/a.png">\n<img width="600" alt="Full decomposition of a sequence" '
        'src="x">\n<img alt="" src="y">\n'
    )
    assert (
        "<img" not in out
        and "[image: Full decomposition of a sequence]" in out
        and "a.png" not in out
    )


def test_figure_img_figcaption_keeps_alt_and_caption_text():
    out = convert_embedded_html(
        '<figure>\n<img src="a.png" alt="Fused module diagram of attention">\n'
        '<figcaption class="c">Fused modules <b>speed up</b> inference</figcaption>\n</figure>\n'
    )
    assert "<fig" not in out and "<img" not in out
    assert (
        "Fused module diagram of attention" in out and "Fused modules **speed up** inference" in out
    )


def test_wrapper_div_spanning_a_fence_is_unwrapped_and_code_kept():
    out = convert_embedded_html(
        '<div class="termy">\n\n```console\n$ fastapi dev\n```\n\n</div>\n\nAfter.'
    )
    assert (
        "<div" not in out
        and "</div>" not in out
        and "$ fastapi dev" in out
        and out.count("```") == 2
    )


def test_raw_html_table_is_converted_to_markdown_table():
    out = convert_embedded_html(
        "<table>\n<tr><th>A</th><th>B</th></tr>\n<tr><td>1</td><td>2</td></tr>\n</table>\n"
    )
    assert "<table" not in out and "| A | B |" in out and "| 1 | 2 |" in out


def test_table_embedded_in_a_pipe_table_cell_is_flattened_to_text():
    row = (
        "| `contact` | `dict` | Info. <details><summary><code>contact</code> "
        "fields</summary><table><thead><tr><th>Parameter</th><th>Type</th></tr>"
        "</thead><tbody><tr><td><code>name</code></td><td><code>str</code></td></tr></tbody>"
        "</table></details> |"
    )
    out = profile_markdown(row + "\n")
    assert "<table" not in out and "<details" not in out and "<td" not in out
    assert (
        "`name` / `str`" in out
        and out.strip().startswith("| `contact`")
        and out.strip().endswith("|")
    )


def test_one_line_div_after_br_is_unwrapped():
    out = convert_embedded_html('    <br><div style="text-align:center">$a+b$</div><br>\n')
    assert "<div" not in out and "<br" not in out and "$a+b$" in out


def test_empty_iframe_and_id_anchor_removed_inside_quote_wrapper():
    out = convert_embedded_html(
        '<iframe width="560" src="https://x"></iframe>\n\n<a id="foo"></a>Hello\n\n> <figure '
        'markdown="span">\n> text\n> </figure>\n'
    )
    assert (
        "<iframe" not in out
        and "<a " not in out
        and "<figure" not in out
        and "Hello" in out
        and "text" in out
    )


def test_html_tags_inside_code_fences_are_untouched():
    src = '```html\n<div class="x"><img src="a"></div>\n```\n'
    assert convert_embedded_html(src).strip() == src.strip()


def test_list_item_keeps_paragraph_code_paragraph_order():
    html = (
        "<ol><li><p>First step:</p><pre>CREATE TABLE a;</pre><p>Then do this:</p><pre>CREATE TABLE "
        "b;</pre>"
        "<p>Finally done.</p></li></ol>"
    )
    md, _ = html_to_markdown(html)
    assert (
        md.index("First step")
        < md.index("CREATE TABLE a")
        < md.index("Then do this")
        < md.index("CREATE TABLE b")
        < md.index("Finally done")
    )
    assert max(len(ln) for ln in md.split("\n")) < 60  # not one flattened line
    assert md.count("```") == 4


def test_definition_with_pre_and_paragraphs_is_not_flattened():
    html = (
        "<dl><dt>Qualified joins</dt><dd><pre>T1 JOIN T2 ON cond\nT1 JOIN T2 USING (c)</pre>"
        "<p>The words INNER and OUTER are optional.</p><p>Second paragraph.</p></dd>"
        "<dt>Short</dt><dd><p>One liner.</p></dd></dl>"
    )
    md, _ = html_to_markdown(html)
    assert (
        "**Qualified joins**:" in md
        and "T1 JOIN T2 ON cond\nT1 JOIN T2 USING (c)" in md
        and "**Short**: One liner." in md
    )
    assert "The words INNER and OUTER are optional.\n\nSecond paragraph." in md


# ---------------------------------------------------------------------------
# rst / sphinx-gallery
# ---------------------------------------------------------------------------


def test_short_tilde_underline_followed_by_text_is_a_heading_and_keeps_fence_parity():
    rst = (
        "Intro text.\n\nGradient Clipping and Optimizer with "
        "DTensor\n~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~\n"
        "**command**: ``torchrun``\n\n.. code-block:: python\n\n    # a comment line\n    x = 1\n\n"
        "Next section\n~~~~~~~~~~~~\n\n.. code-block:: python\n\n    # another comment\n    y = 2\n"
    )
    out = rst_to_markdown(rst)
    assert "~~~" not in out
    assert count_fences(out) % 2 == 0 and count_fences(out) == 4
    assert (
        "# Gradient Clipping and Optimizer with DTensor" in out
        or "## Gradient Clipping and Optimizer with DTensor" in out
    )


def test_gallery_style_sections_inside_module_docstring_become_headings():
    py = (
        '"""\nTitle\n*****\n\nIntro.\n\n#########################\n# Overview - What is it?\n'
        '# ---------------------------------\n\nBody text.\n"""\nimport torch\n'
    )
    out = profile_rst(py, "x.py")
    assert "---------------------------------" not in out
    assert "Overview - What is it?" in out and "\n# Overview" not in out
    assert any(ln.startswith("#") and "Overview - What is it?" in ln for ln in out.split("\n"))


# ---------------------------------------------------------------------------
# junk filter false positives
# ---------------------------------------------------------------------------


def test_name_list_with_a_12_word_intro_sentence_is_not_junk():
    text = (
        "A number of settings can be used to control Django's CSRF behavior:\n\n* "
        "`CSRF_COOKIE_AGE`\n* `CSRF_COOKIE_DOMAIN`\n* `CSRF_COOKIE_NAME`\n* `CSRF_FAILURE_VIEW`\n"
    )
    assert junk_category(text, token_count=100, namespace="technical") is None


def test_bare_name_list_without_prose_is_still_junk():
    text = (
        "* `CSRF_COOKIE_AGE`\n* `CSRF_COOKIE_DOMAIN`\n* `CSRF_COOKIE_NAME`\n* `CSRF_FAILURE_VIEW`\n"
    )
    assert junk_category(text, token_count=100, namespace="technical") == "bare-name-list"


def test_link_list_with_intro_sentence_is_not_junk():
    text = (
        "vLLM supports a variety of methods of speculative decoding. Model-based methods provide "
        "the best latency reduction.\n\n"
        "- [EAGLE](eagle.md)\n- [Draft Model](draft_model.md)\n- [N-Gram](n_gram.md)\n- [Suffix "
        "Decoding](suffix.md)\n"
    )
    assert junk_category(text, token_count=120, namespace="technical") is None


def test_link_list_with_short_intro_is_junk_as_before():
    # cp-algorithms practice-problem lists: heading + bare links, no sentence
    text = "## Practice Problems\n\n- [A](a.md)\n- [B](b.md)\n- [C](c.md)\n- [D](d.md)\n"
    assert junk_category(text, token_count=100, namespace="technical") == "link-index"


def test_link_items_with_descriptive_text_are_not_junk():
    text = (
        "- [Operations](ops.md) - how to run Kafka in production and what to monitor closely\n"
        "- [Security](sec.md) - encryption authentication and authorization setup in detail\n"
        "- [Design](design.md) - the motivation behind the log based architecture overall\n"
    )
    assert junk_category(text, token_count=100, namespace="technical") is None


# ---------------------------------------------------------------------------
# residue scanner
# ---------------------------------------------------------------------------


def test_residue_heading_junk_flags_only_symbol_only_headings():
    assert "heading_junk" in residue.scan_text("# ---------------------------------\n\nbody")
    assert "heading_junk" in residue.scan_text("text", ["Title", ")"])
    for ok in ("A", "C++", "1E + 1P + 1D", "b12x"):
        assert "heading_junk" not in residue.scan_text(f"## {ok}\n\nbody", [ok])


def test_residue_cjk_threshold_ignores_short_gloss_but_flags_translations():
    assert "cjk" not in residue.scan_text("hard to understand (很难理解) why")
    assert "cjk" in residue.scan_text(
        "*[日本語](README-ja.md) ∙ [简体中文](README-zh-Hans.md) 这是一个翻译的页面"
    )


def test_residue_inline_code_wrapping_a_newline_is_not_html_residue():
    assert "html_tag" not in residue.scan_text("the output will be `<p>Joel<br>is a\nslug</p>`.")
    assert "html_tag" in residue.scan_text("raw <div class='x'> leftover")


def test_residue_flags_markers_tabs_triple_colon_and_unbalanced_fence():
    hits = residue.scan_text("--8<-- [start:a]\n:::tip\ntab.selected('x')\n```python\nx\n")
    assert {"snippet_marker", "triple_colon", "tab_selected", "fence_unbalanced"} <= set(hits)


# ---------------------------------------------------------------------------
# end-to-end through the generic cleaner: fences untouched, nothing re-introduced
# ---------------------------------------------------------------------------


def test_profile_plus_generic_clean_keeps_fence_parity():
    src = ':::tip Title\n\nUse it.\n\n```html\n<div class="termy">x</div>\n```\n\n:::\n'
    pre = profile_markdown(src)
    assert count_fences(clean_markdown(pre)) == count_fences(pre) == 2
    assert '<div class="termy">x</div>' in clean_markdown(pre)


# ---------------------------------------------------------------------------
# dropped-chunk persistence
# ---------------------------------------------------------------------------


def _src(**options):
    return Source(
        name="demo",
        namespace="technical",
        priority="P0",
        url="u",
        loader="git",
        license="l",
        auto_level="full",
        trust_level="official",
        refresh_cycle="q",
        domain="d",
        options=options,
    )


def test_dropped_chunks_persisted_with_reason_and_skippable_categories(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "DROPPED_DIR", tmp_path / "dropped")
    rows = [
        {"chunk_id": "a", "text": "x", "reason": "too-short"},
        {"chunk_id": "b", "text": "y", "reason": "link-index"},
    ]
    p = cli.write_dropped_jsonl(_src(), rows)
    got = [json.loads(line) for line in p.read_text().splitlines()]
    assert [(r["chunk_id"], r["reason"]) for r in got] == [("a", "too-short"), ("b", "link-index")]

    p = cli.write_dropped_jsonl(_src(dropped_not_persisted=["link-index"]), rows)
    assert [json.loads(line)["chunk_id"] for line in p.read_text().splitlines()] == ["a"]

    assert (
        cli.write_dropped_jsonl(_src(), []) is None
        and not (tmp_path / "dropped" / "demo.jsonl").exists()
    )


# ---------------------------------------------------------------------------
# crawler: redirects stay in scope, robots group addressed to our UA is binding
# ---------------------------------------------------------------------------


class RedirectNet(Net):
    def __init__(self, pages, redirects, **kw):
        super().__init__(pages, **kw)
        self.redirects = redirects

    def __call__(self, url, **kw):
        resp = super().__call__(url, **kw)
        if url in self.redirects:
            resp = FakeResp(resp.text or "<html></html>")
            resp.url = self.redirects[url]
        return resp


def _crawler(tmp_path, net, **cfg):
    base = {
        "prefix": PREFIX,
        "start_urls": [PREFIX + "index.html"],
        "delay_seconds": 1.0,
        "max_pages": 20,
        "allow": ["ddl*.html"],
        "deny": [],
    }
    base.update(cfg)
    t = {"now": 0.0}

    def sleep(s):
        t["now"] += s

    return Crawler(base, tmp_path, http_get=net, sleep=sleep, clock=lambda: t["now"])


def test_redirect_to_out_of_scope_url_is_not_stored(tmp_path):
    pages = {
        PREFIX + "index.html": page("I", ["ddl.html", "ddl2.html"]),
        PREFIX + "ddl.html": page("D"),
        PREFIX + "ddl2.html": page("D2"),
    }
    net = RedirectNet(pages, {PREFIX + "ddl.html": "https://elsewhere.org/ddl.html"})
    c = _crawler(tmp_path, net)
    got = c.crawl()
    assert [r for _, r, _ in got] == ["ddl2.html"]
    assert c.stats["redirect_out_of_scope"] == 1 and any(
        "redirect out of scope" in e for e in c.errors
    )


def test_redirect_to_robots_disallowed_url_is_not_stored(tmp_path):
    pages = {PREFIX + "index.html": page("I", ["ddl.html"]), PREFIX + "ddl.html": page("D")}
    net = RedirectNet(
        pages,
        {PREFIX + "ddl.html": PREFIX + "ddl-blocked.html"},
        robots="User-agent: *\nDisallow: /docs/current/ddl-blocked.html\n",
    )
    c = _crawler(tmp_path, net, allow=["ddl*.html"])
    assert c.crawl() == [] and c.stats["robots_blocked"] == 1


def test_robots_group_addressed_to_our_user_agent_is_binding_even_if_star_allows(tmp_path):
    pages = {
        PREFIX + "index.html": page("I", ["ddl.html", "ddl2.html"]),
        PREFIX + "ddl.html": page("D"),
        PREFIX + "ddl2.html": page("D2"),
    }
    robots = (
        "User-agent: *\nAllow: /\n\nUser-agent: IntelliHireRAGBot\nDisallow: "
        "/docs/current/ddl.html\n"
    )
    c = _crawler(tmp_path, Net(pages, robots=robots))
    assert [r for _, r, _ in c.crawl()] == ["ddl2.html"] and c.stats["robots_blocked"] == 1


# ---------------------------------------------------------------------------
# xlsx loader: Level 1..6 columns and the 0-role warning
# ---------------------------------------------------------------------------


def _level_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "Sector",
            "Track",
            "Job Role",
            "Job Role Description",
            "Critical Work Function",
            "Key Tasks",
        ]
    )
    ws.append(["ICT", "Data", "Data Engineer", "Builds pipelines.", "Build", "Design ETL"])
    ws2 = wb.create_sheet("TSC")
    ws2.append(
        [
            "Sector",
            "Track",
            "Job Role",
            "TSC Category",
            "TSC Title",
            "Level 1",
            "Level 2",
            "Level 3",
            "Level 4",
            "Level 5",
            "Level 6",
        ]
    )
    ws2.append(
        ["ICT", "Data", "Data Engineer", "Dev", "Data Design", None, None, "x", None, "x", None]
    )
    ws2.append([None, None, None, "Dev", "Data Engineering", None, "x", None, None, None, None])
    wb.save(path)


def test_tsc_level_1_to_6_columns_use_highest_marked_level(tmp_path):
    f = tmp_path / "ict.xlsx"
    _level_workbook(f)
    role = parse_workbook(f)[0]
    assert {(t["title"], t["level"]) for t in role.tsc} == {
        ("Data Design", "Level 5"),
        ("Data Engineering", "Level 2"),
    }


def test_workbook_with_zero_roles_is_reported_loudly(tmp_path, capsys):
    (tmp_path / "skillsfuture-ict").mkdir()
    wb = openpyxl.Workbook()
    wb.active.append(["Nothing", "useful"])
    wb.save(tmp_path / "skillsfuture-ict" / "empty.xlsx")
    src = Source(
        name="skillsfuture-ict",
        namespace="job",
        priority="P0",
        url="https://x",
        loader="xlsx",
        license="l",
        auto_level="semi",
        trust_level="official",
        refresh_cycle="y",
        domain="d",
        options={"parser": "skillsfuture_ict", "glob": "*.xlsx"},
    )
    res = xlsx_fetch(src, tmp_path)
    assert res.skipped and not res.docs
    assert any("0 job roles" in e for e in res.errors) and "0 job roles" in capsys.readouterr().out


def test_inline_svg_is_replaced_by_its_aria_label():
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10" role="img" aria-label="The '
        'big jump of every vertex on a path">'
        '<path d="M0 0L1 1"/><text>1</text></svg>'
    )
    out = convert_embedded_html("Before\n\n" + svg + "\n\nAfter")
    assert (
        "<svg" not in out
        and "<path" not in out
        and "[image: The big jump of every vertex on a path]" in out
    )
    assert "html_tag" in residue.scan_text("x <svg viewBox='0 0 1 1'></svg>")


def test_d2l_bare_labels_in_prose_are_rendered_but_inline_code_is_kept():
    out = profile_d2l(
        "See sec_valueiter and (see subsec_valueitercode); but `sec_keep` and my_sec_x stay.\n"
    )
    assert "the valueiter section" in out and "the valueitercode section" in out
    assert "`sec_keep`" in out and "my_sec_x" in out


def test_d2l_ref_role_target_label_is_rendered_as_text():
    out = profile_d2l("value iteration in :ref:`sec_valueiter` corresponds to it.\n")
    assert "sec_valueiter" not in out and ":ref:" not in out and "the valueiter section" in out
