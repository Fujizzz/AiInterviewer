from pathlib import Path

import pytest

# Optional rag-extra dependencies: skip the whole module when they are not installed.
pytest.importorskip("openpyxl")
import openpyxl

from rag.loaders.xlsx_loader import fetch, parse_workbook, role_to_text, roles_to_docs
from rag.registry import Source


def make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Role Description"
    ws.append(
        [
            "Sector",
            "Track",
            "Job Role",
            "Job Role Description",
            "Critical Work Function",
            "Key Tasks",
        ]
    )
    ws.append(
        [
            "Infocomm Technology",
            "Software and Applications",
            "Software Engineer",
            "Designs and builds software.",
            "Develop software",
            "Write code",
        ]
    )
    ws.append([None, None, None, None, None, "Review code"])
    ws.append([None, None, None, None, "Test software", "Write unit tests"])
    ws.append(
        [
            "Infocomm Technology",
            "Data and AI",
            "Data Engineer",
            "Builds data pipelines.",
            "Build pipelines",
            "Design ETL",
        ]
    )
    ws2 = wb.create_sheet("Job Role_TSC")
    ws2.append(
        ["Sector", "Track", "Job Role", "TSC Category", "TSC Title", "TSC Proficiency Level"]
    )
    ws2.append(
        [
            "Infocomm Technology",
            "Software and Applications",
            "Software Engineer",
            "Development",
            "Software Design",
            "Level 4",
        ]
    )
    ws2.append([None, None, None, "Development", "Programming and Coding", "Level 3"])
    ws3 = wb.create_sheet("Job Role_CCS")
    ws3.append(["Sector", "Track", "Job Role", "Critical Core Skills Title", "Proficiency Level"])
    ws3.append(
        [
            "Infocomm Technology",
            "Software and Applications",
            "Software Engineer",
            "Problem Solving",
            "Intermediate",
        ]
    )
    ws4 = wb.create_sheet("Cover")
    ws4.append(["Skills Framework for ICT"])
    wb.save(path)


def test_parse_workbook_by_headers(tmp_path):
    f = tmp_path / "ict.xlsx"
    make_workbook(f)
    roles = {r.title: r for r in parse_workbook(f)}
    assert set(roles) == {"Software Engineer", "Data Engineer"}
    se = roles["Software Engineer"]
    assert se.track == "Software and Applications" and se.description.startswith("Designs")
    assert se.cwfs["Develop software"] == ["Write code", "Review code"]  # fill-down of merged cells
    assert se.cwfs["Test software"] == ["Write unit tests"]
    assert {(t["title"], t["level"]) for t in se.tsc} == {
        ("Software Design", "Level 4"),
        ("Programming and Coding", "Level 3"),
    }
    assert se.ccs == [{"title": "Problem Solving", "level": "Intermediate"}]


def test_role_text_and_metadata(tmp_path):
    f = tmp_path / "ict.xlsx"
    make_workbook(f)
    r = [x for x in parse_workbook(f) if x.title == "Software Engineer"][0]
    assert "Job role: Software Engineer" in role_to_text(r) and "Software Design" in role_to_text(r)
    doc = roles_to_docs([r], "ict.xlsx", "v1")[0]
    md = doc.extra["metadata"]
    assert md["role_family"] == "Software and Applications"
    assert {"title": "Software Design", "category": "Development", "level": "Level 4"} in md["tsc"]
    assert doc.extra["tsc_titles"] == ["Programming and Coding", "Software Design"]


def src(tmp_path):
    return Source(
        name="skillsfuture-ict",
        namespace="job",
        priority="P0",
        url="https://www.skillsfuture.gov.sg/skills-framework/ict",
        loader="xlsx",
        license="x",
        auto_level="semi",
        trust_level="official",
        refresh_cycle="yearly",
        domain="d",
        options={"parser": "skillsfuture_ict", "glob": "*.xlsx"},
    )


def test_fetch_without_file_reports_manual_download_and_url(tmp_path):
    res = fetch(src(tmp_path), tmp_path)
    assert res.skipped and "MANUAL DOWNLOAD REQUIRED" in res.skip_reason
    assert "https://www.skillsfuture.gov.sg/skills-framework/ict" in res.skip_reason


def test_fetch_reads_xlsx_from_raw_dir(tmp_path):
    (tmp_path / "skillsfuture-ict").mkdir()
    make_workbook(tmp_path / "skillsfuture-ict" / "ict.xlsx")
    res = fetch(src(tmp_path), tmp_path)
    assert not res.skipped and len(res.docs) == 2
    assert res.docs[0].doc_version.startswith("ict.xlsx@sha1:")


def test_corrupt_workbook_is_reported_not_crashing(tmp_path):
    (tmp_path / "skillsfuture-ict").mkdir()
    (tmp_path / "skillsfuture-ict" / "bad.xlsx").write_bytes(b"not an excel file")
    res = fetch(src(tmp_path), tmp_path)
    assert res.skipped and "could not parse" in res.skip_reason


def test_manual_download_note_never_contains_absolute_paths(tmp_path):
    """The note is written into committed pipeline reports: no home/tmp directory may leak."""
    res = fetch(src(tmp_path), tmp_path)
    assert str(tmp_path) not in res.skip_reason and "/Users/" not in res.skip_reason
    repo_raw = Path(__file__).resolve().parents[2] / "rag" / "data" / "raw"
    res = fetch(src(tmp_path), repo_raw)
    assert "<repo>/rag/data/raw/" in res.skip_reason and str(repo_raw) not in res.skip_reason
